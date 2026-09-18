"""Respaldos de consulta: clientes y todas las notas, sin credenciales de usuarios."""
import hashlib
import json
import os
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZIP_DEFLATED, ZipFile
from zoneinfo import ZoneInfo

from django.conf import settings
from django.db import connection, transaction
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .excel_exports import (
    MAX_TEMPLATE_ITEMS,
    build_note_like_cell_writes,
    number_cell,
    render_excel_file,
)
from .models import Client, ClientAddress, Order

BACKUP_TIME_ZONE = ZoneInfo("America/Mexico_City")


@dataclass
class BackupSnapshot:
    captured_at: datetime
    clients: list
    addresses: list
    orders: list


def read_backup_snapshot() -> BackupSnapshot:
    # Una sola instantánea para que Excel, resúmenes y artículos coincidan.
    with transaction.atomic():
        if connection.vendor == "postgresql":
            with connection.cursor() as cursor:
                cursor.execute("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ, READ ONLY")
        return BackupSnapshot(
            captured_at=timezone.now(),
            clients=list(Client.objects.order_by("pk")),
            addresses=list(ClientAddress.objects.select_related("client").order_by("pk")),
            orders=list(Order.objects.select_related("quotation", "quotation__client").prefetch_related(
                "quotation__equipment_items", "extra_costs", "workflow_events",
            ).order_by("pk")),
        )


def append_row(sheet, values):
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            # Conservar teléfonos/folios y evitar fórmulas introducidas como texto.
            cell.data_type = "s"
        elif isinstance(cell.value, Decimal):
            cell.number_format = '#,##0.00;[Red](#,##0.00);"—"'
        elif isinstance(cell.value, datetime):
            cell.number_format = "yyyy-mm-dd hh:mm"
        elif isinstance(cell.value, date):
            cell.number_format = "yyyy-mm-dd"


def local_datetime(value):
    return value.astimezone(BACKUP_TIME_ZONE).replace(tzinfo=None) if value else None


def save_tables(path: Path, tables):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, headers, rows in tables:
        sheet = workbook.create_sheet(title)
        append_row(sheet, headers)
        for row in rows:
            append_row(sheet, row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="1D2F58")
            cell.font = Font(color="FFFFFF", bold=True)
        for index, header in enumerate(headers, 1):
            sheet.column_dimensions[get_column_letter(index)].width = min(40, max(18, len(header) + 3))
    try:
        workbook.save(path)
    finally:
        workbook.close()


def client_tables(snapshot):
    return [
        ("Clientes", ["ID interno", "Código", "Nombre", "Contacto", "Teléfono", "Correo", "Dirección", "Creado", "Actualizado"],
         ((c.pk, c.code, c.client_name, c.contact_person, c.phone_number, c.email, c.address,
           local_datetime(c.created_at), local_datetime(c.updated_at)) for c in snapshot.clients)),
        ("Direcciones", ["ID dirección", "Código cliente", "Dirección", "Colonia", "Referencia"],
         ((a.pk, a.client.code, a.address_line, a.neighborhood, a.reference) for a in snapshot.addresses)),
    ]


def note_rows(orders):
    for order in orders:
        note = order.quotation
        yield (
            order.pk, order.order_id, note.client.code if note.client else "", note.client_name,
            note.phone_number, note.address, note.neighborhood, note.reference, note.delivery_instructions,
            note.delivery_date, note.event_date, note.collection_date,
            order.get_operational_status_display(), order.get_billing_status_display(),
            "Sí" if order.is_cancelled else "No", note.subtotal, note.freight, note.tax_amount,
            note.security_deposit, note.total_estimated, note.discount, note.advance_payment,
            note.total_estimated - note.discount - note.advance_payment,
            order.maps_url, order.assigned_driver_id, "Sí" if order.office_closed else "No",
            local_datetime(order.confirmed_at), local_datetime(order.updated_at),
        )


def note_tables(orders):
    return [
        ("Notas", ["ID interno", "Folio", "Código cliente", "Cliente", "Teléfono", "Dirección", "Colonia",
                   "Referencia", "Instrucciones", "Entrega", "Evento", "Recolección", "Estado operativo",
                   "Estado cobro", "Cancelada", "Subtotal", "Flete", "IVA", "Depósito", "Total",
                   "Descuento", "Anticipo", "Saldo", "Enlace Maps", "ID chofer", "Registro cerrado",
                   "Confirmada", "Actualizada"], note_rows(orders)),
        ("Artículos", ["Folio", "ID artículo", "ID producto", "Cantidad", "Equipo", "Precio unitario", "Importe"],
         ((o.order_id, i.pk, i.inventory_product_id, i.quantity, i.equipment, i.unit_price, i.total)
          for o in orders for i in o.quotation.equipment_items.all())),
        ("Costos extra", ["Folio", "Concepto", "Monto", "Fecha"],
         ((o.order_id, c.concepto, c.monto, local_datetime(c.created_at)) for o in orders for c in o.extra_costs.all())),
        ("Historial", ["Folio", "Categoría", "Estado anterior", "Estado nuevo", "Comentario", "ID usuario", "Fecha"],
         ((o.order_id, e.get_category_display(), e.from_status, e.to_status, e.comment, e.changed_by_id,
           local_datetime(e.created_at)) for o in orders for e in o.workflow_events.all())),
    ]


def write_individual_note(order, path):
    if len(order.quotation.equipment_items.all()) > MAX_TEMPLATE_ITEMS:
        # Una nota heredada larga conserva TODOS sus renglones en formato tabular.
        save_tables(path, note_tables([order]))
        return "tabular"
    writes = build_note_like_cell_writes(
        document_id=order.order_id, created_at=order.confirmed_at.astimezone(BACKUP_TIME_ZONE),
        quotation=order.quotation,
    )
    # Guardar el saldo como valor permite consultarlo sin un motor de recálculo.
    note = order.quotation
    balance = note.total_estimated - note.discount - note.advance_payment
    writes = [number_cell("J49", balance) if cell.cell_reference == "J49" else cell for cell in writes]
    render_excel_file(Path(settings.NOTE_EXCEL_TEMPLATE_PATH), path, writes)
    return "plantilla"


def create_backup_archive(output_dir: Path) -> Path:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    snapshot = read_backup_snapshot()
    stamp = snapshot.captured_at.astimezone(BACKUP_TIME_ZONE).strftime("%Y-%m-%d_%H-%M-%S_%f")
    destination = output_dir / f"Respaldo_Crystal_{stamp}.zip"
    with TemporaryDirectory(prefix="crystal-backup-", dir=output_dir) as temporary:
        root = Path(temporary)
        save_tables(root / "Clientes.xlsx", client_tables(snapshot))
        save_tables(root / "Registro_Notas.xlsx", note_tables(snapshot.orders))
        notes_dir = root / "Notas"
        notes_dir.mkdir()
        index = []
        for order in snapshot.orders:
            folio = re.sub(r"[^A-Za-z0-9_-]", "_", order.order_id) or "nota"
            # El PK impide colisiones de nombres tras sanear folios heredados.
            path = notes_dir / f"{folio}__{order.pk}.xlsx"
            layout = write_individual_note(order, path)
            index.append({"folio": order.order_id, "archivo": path.relative_to(root).as_posix(),
                          "formato": layout, "cancelada": order.is_cancelled,
                          "estado_operativo": order.operational_status, "estado_cobro": order.billing_status})
        files = []
        for path in sorted(root.rglob("*.xlsx")):
            with path.open("rb") as stream:
                checksum = hashlib.file_digest(stream, "sha256").hexdigest()
            files.append({"archivo": path.relative_to(root).as_posix(), "sha256": checksum})
        manifest = {
            "version": 1, "fecha": snapshot.captured_at.astimezone(BACKUP_TIME_ZONE).isoformat(),
            "zona_horaria": str(BACKUP_TIME_ZONE), "clientes": len(snapshot.clients),
            "direcciones": len(snapshot.addresses), "notas": len(snapshot.orders),
            "archivos": files, "indice_notas": index,
            "alcance": "Respaldo de consulta: clientes y notas activas, recogidas y canceladas. No es una copia completa de PostgreSQL; no incluye cotizaciones sin confirmar ni credenciales.",
        }
        (root / "Resumen.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        temporary_zip = root / "backup.zip"
        with ZipFile(temporary_zip, "w", ZIP_DEFLATED) as archive:
            for path in sorted(root.rglob("*")):
                if path.is_file() and path != temporary_zip:
                    archive.write(path, path.relative_to(root).as_posix())
        os.chmod(temporary_zip, 0o600)
        # Publicar solo cuando todos los archivos se generaron correctamente.
        temporary_zip.replace(destination)
    return destination
