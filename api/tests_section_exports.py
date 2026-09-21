from io import BytesIO
from unittest.mock import patch
from zipfile import ZipFile

from django.contrib.auth import get_user_model
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from .models import Client, ClientAddress, Order, OrderExtraCost, set_user_role
from .services import create_order_from_note
from .tests import build_quotation_payload


class SectionExportTests(APITestCase):
    endpoints = (
        ("clients-export-excel", "Clientes.xlsx"),
        ("active-orders-export-excel", "Notas_Activas.xlsx"),
        ("archived-orders-export-excel", "Registro_Notas.xlsx"),
    )

    def setUp(self):
        self.user = get_user_model().objects.create_user(username="export-sales")
        set_user_role(self.user, "ventas")
        self.client.force_authenticate(self.user)

    def export(self, endpoint, filename, params=None):
        with patch("api.excel_exports._generate_pdf") as pdf:
            response = self.client.get(reverse(endpoint), params or {})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response["Content-Type"], "application/zip")
            self.assertIn("attachment;", response["Content-Disposition"])
            self.assertEqual(response["Cache-Control"], "private, no-store")
            content = b"".join(response.streaming_content)
            response.close()
            pdf.assert_not_called()
        with ZipFile(BytesIO(content)) as archive:
            self.assertEqual(archive.namelist(), [filename])
            book = load_workbook(BytesIO(archive.read(filename)))
        self.addCleanup(book.close)
        return book

    def test_all_clients_and_addresses_exported_without_pagination_or_search(self):
        for index in range(32):
            Client.objects.create(client_name=f"Cliente {index}")
        customer = Client.objects.create(client_name='=SUM(1,2)', phone_number="0012345678")
        ClientAddress.objects.create(client=customer, address_line="Dirección alternativa")
        book = self.export(*self.endpoints[0], {"page": 2, "search": "inexistente"})
        self.assertEqual(book["Clientes"].max_row, 34)
        self.assertEqual(book["Clientes"]["C34"].value, "=SUM(1,2)")
        self.assertEqual(book["Clientes"]["C34"].data_type, "s")
        self.assertEqual(book["Clientes"]["E34"].value, "0012345678")
        self.assertEqual(book["Direcciones"]["C2"].value, "Dirección alternativa")
        self.assertEqual(book["Clientes"].freeze_panes, "A2")

    def test_active_and_archive_match_section_membership_and_preserve_details(self):
        active = create_order_from_note(build_quotation_payload())
        cancelled = create_order_from_note(build_quotation_payload())
        cancelled.is_cancelled = True
        cancelled.save()
        archived = create_order_from_note(build_quotation_payload())
        archived.operational_status = Order.OperationalStatus.RECOGIDO
        archived.save()
        OrderExtraCost.objects.create(order=active, concepto="Maniobras", monto="125.50")
        before = list(Order.objects.values())
        book = self.export(*self.endpoints[1], {"folder": "por-cobrar", "page": 10})
        self.assertEqual([row[1].value for row in list(book["Notas"].rows)[1:]],
                         [active.order_id, cancelled.order_id])
        self.assertEqual(book["Notas"]["O3"].value, "Sí")
        self.assertEqual(book["Notas"]["W2"].value, 3650)
        self.assertEqual(book["Artículos"].max_row, 3)
        self.assertEqual(book["Costos extra"]["C2"].value, 125.50)
        archive = self.export(*self.endpoints[2])
        self.assertEqual(archive["Notas"].max_row, 2)
        self.assertEqual(archive["Notas"]["B2"].value, archived.order_id)
        self.assertEqual(before, list(Order.objects.values()))

    def test_empty_sections_produce_readable_workbooks_with_headers(self):
        for endpoint, filename in self.endpoints:
            with self.subTest(endpoint=endpoint):
                book = self.export(endpoint, filename)
                self.assertTrue(all(sheet.max_row == 1 for sheet in book))

    def test_anonymous_and_drivers_cannot_export(self):
        for role in (None, "chofer"):
            if role:
                set_user_role(self.user, role)
                self.user.refresh_from_db()
            self.client.force_authenticate(self.user if role else None)
            for endpoint, _ in self.endpoints:
                with self.subTest(role=role, endpoint=endpoint):
                    self.assertIn(self.client.get(reverse(endpoint)).status_code, (401, 403))

    def test_admin_can_export_every_section(self):
        set_user_role(self.user, "admin")
        self.user.refresh_from_db()
        for endpoint, filename in self.endpoints:
            self.export(endpoint, filename)
