import hashlib
import json
from contextlib import contextmanager
from datetime import datetime, timezone
from io import BytesIO, StringIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import MagicMock, patch
from zipfile import ZipFile

from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import SimpleTestCase, TestCase
from openpyxl import Workbook, load_workbook

from .backups import create_backup_archive
from .drive_backup import DriveBackupError, DriveBackupUploader
from .models import ClientAddress, Order, OrderExtraCost, QuotationItem
from .services import create_order_from_note
from .tests import build_quotation_payload


class BackupArchiveTests(TestCase):
    @contextmanager
    def workspace(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            template = root / "template.xlsx"
            Workbook().save(template)
            with self.settings(NOTE_EXCEL_TEMPLATE_PATH=template):
                yield root / "output"

    def make_order(self):
        return create_order_from_note(build_quotation_payload())

    def test_all_statuses_clients_addresses_items_and_money_are_exported(self):
        active = self.make_order()
        archived = self.make_order()
        archived.operational_status = Order.OperationalStatus.RECOGIDO
        archived.save()
        cancelled = self.make_order()
        cancelled.is_cancelled = True
        cancelled.save()
        client = active.quotation.client
        client.client_name = '=HYPERLINK("https://invalid.example","Cliente")'
        client.phone_number = "0012345678"
        client.save()
        ClientAddress.objects.create(client=client, address_line="Dirección alternativa")
        OrderExtraCost.objects.create(order=active, concepto="Maniobras", monto="125.50")
        before = list(Order.objects.values())
        with self.workspace() as output, patch("api.excel_exports._generate_pdf") as pdf:
            path = create_backup_archive(output)
            pdf.assert_not_called()
            with ZipFile(path) as archive:
                metadata = json.loads(archive.read("Resumen.json"))
                self.assertEqual(metadata["notas"], 3)
                self.assertEqual(metadata["direcciones"], 1)
                self.assertEqual(len([name for name in archive.namelist() if name.startswith("Notas/")]), 3)
                for entry in metadata["archivos"]:
                    self.assertEqual(hashlib.sha256(archive.read(entry["archivo"])).hexdigest(), entry["sha256"])
                clients = load_workbook(BytesIO(archive.read("Clientes.xlsx")))
                self.assertEqual(clients["Clientes"]["C2"].data_type, "s")
                self.assertEqual(clients["Clientes"]["E2"].value, "0012345678")
                self.assertEqual(clients["Direcciones"]["C2"].value, "Dirección alternativa")
                notes = load_workbook(BytesIO(archive.read("Registro_Notas.xlsx")))
                self.assertEqual(notes["Notas"].max_row, 4)
                self.assertEqual(notes["Notas"]["O4"].value, "Sí")
                self.assertEqual(notes["Notas"]["W2"].value, 3650)
                self.assertEqual(notes["Artículos"]["F2"].value, 85)
                self.assertEqual(notes["Costos extra"]["C2"].value, 125.50)
                note_path = metadata["indice_notas"][0]["archivo"]
                note = load_workbook(BytesIO(archive.read(note_path)), data_only=True)
                self.assertEqual(note.active["J49"].value, 3650)
                note.close()
                clients.close()
                notes.close()
        self.assertEqual(before, list(Order.objects.values()))

    def test_real_template_can_be_opened_with_balance_without_recalculation(self):
        self.make_order()
        template = Path(__file__).resolve().parent.parent / "templates/excel/Nota.xlsx"
        with TemporaryDirectory() as temporary, self.settings(NOTE_EXCEL_TEMPLATE_PATH=template):
            with ZipFile(create_backup_archive(Path(temporary))) as archive:
                entry = json.loads(archive.read("Resumen.json"))["indice_notas"][0]
                book = load_workbook(BytesIO(archive.read(entry["archivo"])), data_only=True)
                self.assertEqual(book.active["J49"].value, 3650)
                book.close()

    def test_sanitized_folios_cannot_escape_or_overwrite_other_notes(self):
        first, second = self.make_order(), self.make_order()
        Order.objects.filter(pk=first.pk).update(order_id="../one")
        Order.objects.filter(pk=second.pk).update(order_id="..\\one")
        with self.workspace() as output:
            with ZipFile(create_backup_archive(output)) as archive:
                names = [name for name in archive.namelist() if name.startswith("Notas/")]
                self.assertEqual(len(set(names)), 2)
                self.assertTrue(all(".." not in name and "\\" not in name for name in names))

    def test_long_legacy_note_preserves_every_item(self):
        order = self.make_order()
        for index in range(22):
            QuotationItem.objects.create(quotation=order.quotation, equipment=f"Extra {index}", quantity=1)
        with self.workspace() as output:
            with ZipFile(create_backup_archive(output)) as archive:
                entry = json.loads(archive.read("Resumen.json"))["indice_notas"][0]
                self.assertEqual(entry["formato"], "tabular")
                book = load_workbook(BytesIO(archive.read(entry["archivo"])))
                self.assertEqual(book["Artículos"].max_row, 24)
                book.close()

    def test_failure_never_publishes_partial_archive_or_deletes_older_backup(self):
        self.make_order()
        with self.workspace() as output:
            previous = create_backup_archive(output)
            with patch("api.backups.write_individual_note", side_effect=ValueError("failure")):
                with self.assertRaises(ValueError):
                    create_backup_archive(output)
            self.assertEqual(list(output.iterdir()), [previous])

    def test_filename_uses_mexico_city_date_not_utc_date(self):
        instant = datetime(2026, 9, 18, 0, 0, tzinfo=timezone.utc)
        with self.workspace() as output, patch("api.backups.timezone.now", return_value=instant):
            path = create_backup_archive(output)
            self.assertTrue(path.name.startswith("Respaldo_Crystal_2026-09-17_18-00-00"))

    def test_local_command_does_not_connect_to_drive(self):
        with self.workspace() as output, patch("api.drive_backup.requests.Session") as session:
            call_command("backup_to_drive", local_only=True, output_dir=output, stdout=StringIO())
            self.assertEqual(len(list(output.glob("*.zip"))), 1)
            session.assert_not_called()


def response(payload=None, status=200, headers=None):
    item = MagicMock(status_code=status, headers=headers or {})
    item.json.return_value = payload or {}
    item.__enter__.return_value = item
    return item


class DriveUploadTests(SimpleTestCase):
    credentials = {"client_id": "id", "client_secret": "private", "refresh_token": "private-refresh"}

    def uploader(self, responses):
        session = MagicMock()
        session.request.side_effect = responses
        return DriveBackupUploader("folder-id", self.credentials, session), session

    def folder_response(self):
        return response({"mimeType": "application/vnd.google-apps.folder", "capabilities": {"canAddChildren": True}})

    def test_upload_uses_only_target_folder_and_checks_remote_integrity(self):
        payload = b"example-backup"
        uploader, session = self.uploader([
            response({"access_token": "token"}), self.folder_response(), response({"permissions": [{"type": "user"}]}),
            response(headers={"Location": "https://www.googleapis.com/upload/session"}),
            response({"id": "saved", "size": str(len(payload)), "md5Checksum": hashlib.md5(payload).hexdigest()}),
        ])
        with TemporaryDirectory() as temporary:
            path = Path(temporary) / "backup.zip"
            path.write_bytes(payload)
            self.assertEqual(uploader.upload(path), "saved")
        upload_call = session.request.call_args_list[3]
        self.assertEqual(upload_call.kwargs["json"]["parents"], ["folder-id"])
        self.assertFalse(any(call.args[0] == "DELETE" for call in session.request.call_args_list))
        self.assertTrue(all(call.kwargs["allow_redirects"] is False for call in session.request.call_args_list))

    def test_public_folder_is_rejected_before_any_upload(self):
        uploader, session = self.uploader([
            response({"access_token": "token"}), self.folder_response(), response({"permissions": [{"type": "anyone"}]}),
        ])
        with self.assertRaisesMessage(DriveBackupError, "restringida"):
            uploader.upload(Path("unused.zip"))
        self.assertEqual(session.request.call_count, 3)

    def test_token_errors_do_not_expose_credentials_or_provider_body(self):
        uploader, _ = self.uploader([response({"error": "private-refresh"}, status=401)])
        with self.assertRaises(DriveBackupError) as caught:
            uploader.authenticate()
        self.assertNotIn("private", str(caught.exception))

    def test_missing_google_credentials_fails_before_reading_database(self):
        with patch.dict("os.environ", {}, clear=True), patch("api.management.commands.backup_to_drive.create_backup_archive") as create:
            with self.assertRaises(CommandError):
                call_command("backup_to_drive")
            create.assert_not_called()
