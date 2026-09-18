from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import MagicMock, patch
from zipfile import ZIP_DEFLATED, ZipFile

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.test import SimpleTestCase, override_settings
from django.urls import reverse
from openpyxl import Workbook, load_workbook
from rest_framework.authtoken.models import Token
from rest_framework.test import APITestCase

from .excel_exports import DocumentCellWrite, _apply_cell_writes, render_document_bundle
from .models import UserProfile, set_user_role
from .note_import import NoteImportError, read_note_excel
from .route_optimization import resolve_maps_url_coordinates
from .tests import build_quotation_payload


class AccessSecurityTests(APITestCase):
    def setUp(self):
        cache.clear()
        self.addCleanup(cache.clear)
        self.admin = get_user_model().objects.create_user(username="owner", is_staff=True)
        self.driver = get_user_model().objects.create_user(username="driver")
        set_user_role(self.driver, UserProfile.Role.CHOFER)

    @override_settings(REGISTRATION_ACCESS_KEY="InvitationSecret")
    def test_registration_cannot_create_admin(self):
        response = self.client.post(reverse("api-register"), {
            "email": "new@example.com", "password": "Roble!9274Nube",
            "registrationKey": "InvitationSecret", "role": "admin",
        })
        self.assertEqual(response.status_code, 400)
        self.assertFalse(get_user_model().objects.filter(email="new@example.com").exists())

    @override_settings(REGISTRATION_ACCESS_KEY="")
    def test_registration_is_disabled_without_key(self):
        response = self.client.post(reverse("api-register"), {
            "email": "new@example.com", "password": "Roble!9274Nube",
            "registrationKey": "CrystalRegister2026",
        })
        self.assertEqual(response.status_code, 400)
        self.assertIn("registrationKey", response.data)

    @override_settings(REGISTRATION_ACCESS_KEY="InvitationSecret")
    def test_registration_rejects_common_password(self):
        response = self.client.post(reverse("api-register"), {
            "email": "new@example.com", "password": "123456789",
            "registrationKey": "InvitationSecret",
        })
        self.assertEqual(response.status_code, 400)
        self.assertIn("password", response.data)

    def test_login_throttles_repeated_attempts(self):
        for _ in range(10):
            self.client.post(reverse("api-login"), {"identifier": "missing", "password": "invalid"})
        response = self.client.post(reverse("api-login"), {"identifier": "missing", "password": "invalid"})
        self.assertEqual(response.status_code, 429)

    def test_driver_cannot_bypass_permissions_through_general_endpoints(self):
        self.client.force_authenticate(self.driver)
        for name, method, kwargs in [
            ("order-bulk-status-update", "post", {}),
            ("order-list", "post", {}),
            ("order-detail", "put", {"order_id": "0001-26"}),
            ("quotation-list", "post", {}),
            ("client-list", "get", {}),
            ("accounting-overview", "get", {}),
        ]:
            with self.subTest(name=name):
                response = getattr(self.client, method)(reverse(name, kwargs=kwargs), {}, format="json")
                self.assertEqual(response.status_code, 403)

    def test_driver_can_update_only_assigned_order(self):
        self.client.force_authenticate(self.admin)
        created = self.client.post(reverse("order-list"), build_quotation_payload(), format="json")
        self.assertEqual(created.status_code, 201)
        from .models import Order

        order = Order.objects.get(order_id=created.data["orderId"])
        self.client.force_authenticate(self.driver)
        url = reverse("order-status-update", kwargs={"order_id": order.order_id})
        self.assertEqual(self.client.post(url, {"operationalStatus": "ENTREGADO"}).status_code, 403)
        order.assigned_driver = self.driver
        order.save(update_fields=["assigned_driver"])
        self.assertEqual(self.client.post(url, {"billingStatus": "COBRADO"}).status_code, 403)
        self.assertEqual(self.client.post(url, {"operationalStatus": "ENTREGADO"}).status_code, 200)

    def test_password_change_revokes_token(self):
        token = Token.objects.create(user=self.driver)
        self.client.force_authenticate(self.admin)
        response = self.client.patch(reverse("team-detail", kwargs={"user_id": self.driver.pk}),
                                     {"password": "Roble!9274Nube"})
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Token.objects.filter(key=token.key).exists())

    def test_rejected_role_change_does_not_save_other_changes(self):
        self.client.force_authenticate(self.admin)
        response = self.client.patch(reverse("team-detail", kwargs={"user_id": self.admin.pk}),
                                     {"displayName": "Changed", "role": "ventas"})
        self.assertEqual(response.status_code, 400)
        self.admin.refresh_from_db()
        self.assertEqual(self.admin.first_name, "")


class DocumentAndUrlSecurityTests(SimpleTestCase):
    def setUp(self):
        cache.clear()
        self.addCleanup(cache.clear)

    @patch("api.route_optimization.requests.get")
    def test_untrusted_maps_urls_never_make_requests(self, get):
        for url in ["http://127.0.0.1", "https://169.254.169.254/", "file:///etc/passwd",
                    "https://google.com.evil.test/maps", "https://user@google.com/maps",
                    "https://www.google.com:8443/maps"]:
            self.assertIsNone(resolve_maps_url_coordinates(url))
        get.assert_not_called()

    @patch("api.route_optimization.requests.get")
    def test_maps_redirect_cannot_reach_internal_host(self, get):
        get.return_value = MagicMock(status_code=302, headers={"Location": "http://127.0.0.1/admin"})
        self.assertIsNone(resolve_maps_url_coordinates("https://maps.app.goo.gl/example"))
        self.assertEqual(get.call_count, 1)
        self.assertFalse(get.call_args.kwargs["allow_redirects"])
        get.return_value.close.assert_called_once()

    @patch("api.route_optimization.requests.get")
    def test_safe_redirect_resolves_coordinates(self, get):
        target = "https://www.google.com/maps/@19.05,-98.20,15z"
        get.side_effect = [MagicMock(status_code=302, headers={"Location": target}),
                           MagicMock(status_code=200, url=target)]
        self.assertEqual(resolve_maps_url_coordinates("https://maps.app.goo.gl/example"), (19.05, -98.20))

    def test_export_preserves_user_input_as_text_and_template_formulas(self):
        workbook = Workbook()
        _apply_cell_writes(workbook.active, [
            DocumentCellWrite("A1", "text", '=HYPERLINK("https://evil.test","Click")'),
            DocumentCellWrite("A2", "formula", "=SUM(1,2)"),
        ])
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        saved = load_workbook(output)
        self.assertEqual(saved.active["A1"].data_type, "s")
        self.assertEqual(saved.active["A2"].data_type, "f")
        saved.close()

    def test_import_rejects_oversized_decompressed_archive(self):
        output = BytesIO()
        with ZipFile(output, "w", ZIP_DEFLATED) as archive:
            archive.writestr("large.xml", b"a" * (25 * 1024 * 1024 + 1))
        output.seek(0)
        with self.assertRaises(NoteImportError):
            read_note_excel(output)

    @patch("api.excel_exports._generate_pdf")
    def test_export_folio_cannot_escape_temporary_directory(self, generate_pdf):
        generate_pdf.side_effect = lambda excel, pdf, **kwargs: pdf.write_bytes(b"pdf")
        with TemporaryDirectory() as folder:
            template = Path(folder) / "template.xlsx"
            Workbook().save(template)
            bundle = render_document_bundle(
                template_path=template, output_stem="../../escape'\"", cell_writes=[],
            )
            self.assertEqual(Path(bundle.excel_filename).name, bundle.excel_filename)
            self.assertNotIn("..", bundle.excel_filename)
            self.assertNotIn("'", bundle.excel_filename)
            self.assertNotIn('"', bundle.excel_filename)
            self.assertEqual(bundle.pdf_bytes, b"pdf")
